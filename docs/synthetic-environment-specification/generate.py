"""Generate a portable synthetic environment specification and engineering drawings."""

import argparse
import hashlib
import json
import math
import shutil
import subprocess
import tempfile
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from itertools import pairwise
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin

type JsonObject = dict[str, Any]
type Point2 = tuple[float, float]
type Point3 = tuple[float, float, float]
type PixelPoint = tuple[float, float]

_BUNDLE_DIR = Path(__file__).resolve().parent
_ROOT = Path(__file__).resolve().parents[2]
_DEFAULT_CONFIG = _ROOT / "examples" / "simulator.v2.json"
_DEFAULT_OUTPUT = _BUNDLE_DIR
_DEFAULT_ASSETS = _BUNDLE_DIR / "assets"
_DOCUMENT_BASENAME = "synthetic-environment-specification"
_PLAN_FILENAME = "plan-view.png"
_ISOMETRIC_FILENAME = "isometric-view.png"
_MARKDOWN_FILENAME = "synthetic-environment-specification.md"
_SOURCE_FILENAME = "SOURCE.json"
_FONT_PATH = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
_FIXED_DOCUMENT_TIME = datetime(2000, 1, 1, tzinfo=UTC)
_DRAWING_SCALE = 2
DEFAULT_ANGLE_OF_REPOSE_DEG = 35.0
DEFAULT_SLOPE_RELAXATION_MAX_ITERATIONS = 32

_INK = "#1f2937"
_MUTED = "#64748b"
_GRID = "#dbe4ee"
_BOUNDARY = "#2563eb"
_BOUNDARY_FILL = "#dbeafe"
_SENSOR = "#d97706"
_INLET = "#dc2626"
_RAY = "#059669"
_SENSOR_ONE = "#1976d2"
_SENSOR_TWO = "#009688"
_WHITE = "#ffffff"


class _ObjectView:
    """Expose trusted, Rust-validated public JSON fields to the document renderer."""

    def __init__(self, document: JsonObject) -> None:
        self._document = document

    def __getattr__(self, name: str) -> Any:
        try:
            return _as_view(self._document[name])
        except KeyError as error:
            raise AttributeError(name) from error


def _as_view(value: Any) -> Any:
    if isinstance(value, dict):
        return _ObjectView(value)
    if isinstance(value, list):
        return tuple(_as_view(item) for item in value)
    return value


@dataclass(frozen=True, slots=True)
class PublicInputs:
    """Public JSON values needed by the document renderer."""

    simulator: _ObjectView
    environment: _ObjectView


@dataclass(frozen=True, slots=True)
class SpecificationSources:
    """Validated public inputs for the synthetic environment specification."""

    simulator_path: Path
    environment_path: Path
    quality_path: Path
    inputs: PublicInputs
    simulator_document: JsonObject
    environment_document: JsonObject
    quality_document: JsonObject
    hashes: dict[str, str]
    input_fingerprint: str
    document_generator_sha256: str
    model_values: dict[str, float | int]
    fingerprint: str


@dataclass(frozen=True, slots=True)
class DrawingPaths:
    """Generated engineering drawing paths."""

    plan: Path
    isometric: Path


class _ScaledDraw:
    """Render logical drawing coordinates to a high-resolution raster."""

    def __init__(self, image: Image.Image) -> None:
        self._draw = ImageDraw.Draw(image)

    def line(
        self,
        xy: Any,
        fill: Any = None,
        width: int = 0,
        joint: str | None = None,
    ) -> None:
        self._draw.line(
            _scale_coordinates(xy),
            fill=fill,
            width=width * _DRAWING_SCALE,
            joint=joint,
        )

    def polygon(
        self,
        xy: Any,
        fill: Any = None,
        outline: Any = None,
        width: int = 1,
    ) -> None:
        self._draw.polygon(
            _scale_coordinates(xy),
            fill=fill,
            outline=outline,
            width=width * _DRAWING_SCALE,
        )

    def ellipse(
        self,
        xy: Any,
        fill: Any = None,
        outline: Any = None,
        width: int = 1,
    ) -> None:
        self._draw.ellipse(
            _scale_coordinates(xy),
            fill=fill,
            outline=outline,
            width=width * _DRAWING_SCALE,
        )

    def rectangle(
        self,
        xy: Any,
        fill: Any = None,
        outline: Any = None,
        width: int = 1,
    ) -> None:
        self._draw.rectangle(
            _scale_coordinates(xy),
            fill=fill,
            outline=outline,
            width=width * _DRAWING_SCALE,
        )

    def rounded_rectangle(
        self,
        xy: Any,
        radius: int = 0,
        fill: Any = None,
        outline: Any = None,
        width: int = 1,
    ) -> None:
        self._draw.rounded_rectangle(
            _scale_coordinates(xy),
            radius=radius * _DRAWING_SCALE,
            fill=fill,
            outline=outline,
            width=width * _DRAWING_SCALE,
        )

    def text(
        self,
        xy: PixelPoint,
        text: str,
        *,
        fill: Any = None,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont | None = None,
        anchor: str | None = None,
    ) -> None:
        self._draw.text(
            _scale_coordinates(xy),
            text,
            fill=fill,
            font=font,
            anchor=anchor,
        )

    def multiline_text(
        self,
        xy: PixelPoint,
        text: str,
        *,
        fill: Any = None,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont | None = None,
        anchor: str | None = None,
        spacing: int = 4,
    ) -> None:
        self._draw.multiline_text(
            _scale_coordinates(xy),
            text,
            fill=fill,
            font=font,
            anchor=anchor,
            spacing=spacing * _DRAWING_SCALE,
        )

    def multiline_textbbox(
        self,
        xy: PixelPoint,
        text: str,
        *,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont | None = None,
        anchor: str | None = None,
        spacing: int = 4,
    ) -> tuple[float, float, float, float]:
        box = self._draw.multiline_textbbox(
            _scale_coordinates(xy),
            text,
            font=font,
            anchor=anchor,
            spacing=spacing * _DRAWING_SCALE,
        )
        return tuple(value / _DRAWING_SCALE for value in box)  # type: ignore[return-value]


def _scale_coordinates(value: Any) -> Any:
    if isinstance(value, int | float):
        return value * _DRAWING_SCALE
    if isinstance(value, tuple):
        return tuple(_scale_coordinates(item) for item in value)
    if isinstance(value, list):
        return [_scale_coordinates(item) for item in value]
    return value


def load_sources(simulator_path: Path) -> SpecificationSources:
    """Load and validate the three public JSON sources."""
    simulator_path = simulator_path.resolve()
    if _DEFAULT_CONFIG.is_symlink() or simulator_path != _DEFAULT_CONFIG.absolute():
        raise ValueError("specification simulator config must be examples/simulator.v2.json")
    simulator_document = _load_json(simulator_path)
    expected_references = (
        ("environment_path", _ROOT / "examples" / "environment.v1.json"),
        ("quality_profile_path", _ROOT / "examples" / "quality-profile.v1.json"),
    )
    for field, expected_path in expected_references:
        value = simulator_document.get(field)
        if not isinstance(value, str):
            raise ValueError(f"specification simulator field must be a path string: {field}")
        referenced_path = (simulator_path.parent / value).resolve()
        if expected_path.is_symlink() or referenced_path != expected_path.absolute():
            raise ValueError(f"specification requires {_relative_path(expected_path)}")
    environment_path = expected_references[0][1].resolve()
    quality_path = expected_references[1][1].resolve()
    paths = (environment_path, simulator_path, quality_path)
    documents = (_load_json(environment_path), simulator_document, _load_json(quality_path))
    inputs = PublicInputs(
        simulator=_ObjectView(simulator_document),
        environment=_ObjectView(documents[0]),
    )
    relative_names = tuple(_relative_path(path) for path in paths)
    hashes = {
        name: hashlib.sha256(path.read_bytes()).hexdigest()
        for name, path in zip(relative_names, paths, strict=True)
    }
    input_digest = hashlib.sha256()
    for name in sorted(hashes):
        input_digest.update(name.encode("utf-8"))
        input_digest.update(b"\0")
        input_digest.update(bytes.fromhex(hashes[name]))
    input_fingerprint = input_digest.hexdigest()
    document_generator_sha256 = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    model_values: dict[str, float | int] = {
        "angle_of_repose_deg": DEFAULT_ANGLE_OF_REPOSE_DEG,
        "slope_relaxation_max_iterations": DEFAULT_SLOPE_RELAXATION_MAX_ITERATIONS,
    }
    artifact_digest = hashlib.sha256()
    artifact_digest.update(b"synthetic-environment-artifacts-v1\0")
    artifact_digest.update(bytes.fromhex(input_fingerprint))
    artifact_digest.update(bytes.fromhex(document_generator_sha256))
    artifact_digest.update(
        json.dumps(model_values, allow_nan=False, sort_keys=True, separators=(",", ":")).encode()
    )
    return SpecificationSources(
        simulator_path=simulator_path,
        environment_path=environment_path,
        quality_path=quality_path,
        inputs=inputs,
        simulator_document=documents[1],
        environment_document=documents[0],
        quality_document=documents[2],
        hashes=hashes,
        input_fingerprint=input_fingerprint,
        document_generator_sha256=document_generator_sha256,
        model_values=model_values,
        fingerprint=artifact_digest.hexdigest(),
    )


def polygon_area(boundary: Sequence[Point2]) -> float:
    """Return the unsigned area of a simple polygon."""
    return abs(
        sum(
            first[0] * second[1] - second[0] * first[1]
            for first, second in _polygon_edges(boundary)
        )
        / 2.0
    )


def zero_degree_floor_hit(sensor: JsonObject, floor_z_m: float) -> Point3:
    """Return the floor intersection of a sensor's zero-degree ray."""
    p0 = tuple(float(value) for value in sensor["p0_m"])
    u0 = tuple(float(value) for value in sensor["u0"])
    scale = (floor_z_m - p0[2]) / u0[2]
    return tuple(p0[index] + scale * u0[index] for index in range(3))  # type: ignore[return-value]


def sensor_tilt_from_down_deg(sensor: JsonObject) -> float:
    """Return the zero-ray tilt from the downward vertical axis."""
    u0 = tuple(float(value) for value in sensor["u0"])
    return math.degrees(math.acos(min(1.0, max(-1.0, -u0[2]))))


def sensor_rotation_axis(sensor: JsonObject) -> Point3:
    """Return the right-handed rotation axis u0 cross u90."""
    u0 = tuple(float(value) for value in sensor["u0"])
    u90 = tuple(float(value) for value in sensor["u90"])
    return (
        u0[1] * u90[2] - u0[2] * u90[1],
        u0[2] * u90[0] - u0[0] * u90[2],
        u0[0] * u90[1] - u0[1] * u90[0],
    )


def floor_measurement_segments(
    sensor: JsonObject,
    boundary: Sequence[Point2],
    floor_z_m: float,
) -> tuple[tuple[Point3, Point3], ...]:
    """Clip the scan-plane and floor-plane intersection to the floor polygon."""
    hit = zero_degree_floor_hit(sensor, floor_z_m)
    u90 = tuple(float(value) for value in sensor["u90"])
    line_origin = (hit[0], hit[1])
    line_direction = (u90[0], u90[1])
    if math.hypot(*line_direction) <= 1e-12:
        raise ValueError("sensor u90 must have a horizontal component")

    intersections: list[float] = []
    for first, second in _polygon_edges(boundary):
        edge = (second[0] - first[0], second[1] - first[1])
        delta = (first[0] - line_origin[0], first[1] - line_origin[1])
        denominator = _cross2(line_direction, edge)
        if abs(denominator) <= 1e-12:
            if abs(_cross2(delta, line_direction)) <= 1e-12:
                intersections.extend(
                    _line_parameter(line_origin, line_direction, point) for point in (first, second)
                )
            continue
        edge_parameter = _cross2(delta, line_direction) / denominator
        if -1e-12 <= edge_parameter <= 1.0 + 1e-12:
            intersections.append(_cross2(delta, edge) / denominator)

    parameters = _unique_sorted(intersections)
    segments = []
    for start, end in pairwise(parameters):
        midpoint = _line_point(line_origin, line_direction, (start + end) / 2.0)
        if not _point_in_polygon(midpoint, boundary):
            continue
        start_xy = _line_point(line_origin, line_direction, start)
        end_xy = _line_point(line_origin, line_direction, end)
        segments.append(
            (
                (start_xy[0], start_xy[1], floor_z_m),
                (end_xy[0], end_xy[1], floor_z_m),
            )
        )
    return tuple(segments)


def _cross2(first: Point2, second: Point2) -> float:
    return first[0] * second[1] - first[1] * second[0]


def _line_parameter(origin: Point2, direction: Point2, point: Point2) -> float:
    denominator = direction[0] ** 2 + direction[1] ** 2
    return (
        (point[0] - origin[0]) * direction[0] + (point[1] - origin[1]) * direction[1]
    ) / denominator


def _line_point(origin: Point2, direction: Point2, parameter: float) -> Point2:
    return (
        origin[0] + parameter * direction[0],
        origin[1] + parameter * direction[1],
    )


def _unique_sorted(values: Iterable[float]) -> tuple[float, ...]:
    result: list[float] = []
    for value in sorted(values):
        if not result or not math.isclose(value, result[-1], abs_tol=1e-9):
            result.append(value)
    return tuple(result)


def _point_in_polygon(point: Point2, boundary: Sequence[Point2]) -> bool:
    inside = False
    for first, second in _polygon_edges(boundary):
        if (first[1] > point[1]) == (second[1] > point[1]):
            continue
        intersection_x = (second[0] - first[0]) * (point[1] - first[1]) / (
            second[1] - first[1]
        ) + first[0]
        if point[0] < intersection_x:
            inside = not inside
    return inside


def _polygon_edges(boundary: Sequence[Point2]) -> Iterator[tuple[Point2, Point2]]:
    if not boundary:
        return iter(())
    return pairwise((*boundary, boundary[0]))


def generate_drawings(sources: SpecificationSources, output_dir: Path) -> DrawingPaths:
    """Generate exact plan and isometric environment drawings."""
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = DrawingPaths(
        plan=output_dir / _PLAN_FILENAME,
        isometric=output_dir / _ISOMETRIC_FILENAME,
    )
    _draw_plan(sources, paths.plan)
    _draw_isometric(sources, paths.isometric)
    return paths


def generate_docx(
    sources: SpecificationSources,
    drawings: DrawingPaths,
    output_path: Path,
) -> None:
    """Generate one deterministic DOCX specification from validated inputs."""
    document = Document()
    _configure_document(document, sources)
    _write_title_page(document, sources)
    _write_source_section(document, sources)
    _write_coordinate_section(document, sources)
    _write_space_section(document, sources, drawings)
    _write_sensor_section(document, sources, drawings)
    _write_simulation_section(document, sources)
    _write_measurement_section(document, sources)
    _write_interpretation_section(document, sources)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="synthetic-environment-docx-") as temporary:
        raw_path = Path(temporary) / output_path.name
        document.save(str(raw_path))
        _normalize_docx(raw_path, output_path)


def generate_markdown(sources: SpecificationSources, output_path: Path) -> None:
    """Generate the self-contained human-readable environment specification."""
    environment = sources.inputs.environment
    boundary = _boundary(sources)
    minimum_x, maximum_x, minimum_y, maximum_y = _bounds(boundary)
    area = polygon_area(boundary)
    height = environment.top_z_m - environment.floor_z_m
    lines = [
        "# 공개 합성 LiDAR 시뮬레이션 환경 규격",
        "",
        "이 문서는 합성 통합 검증에 사용하는 공간, 센서 설치, 시나리오와 측정 설정의 "
        "자기완결 전달 문서다. 실제 현장 정보는 포함하지 않는다. 수치 정본은 문서 하단에 "
        "수록한 versioned JSON 3개다.",
        "",
        f"Input fingerprint SHA-256: `{sources.input_fingerprint}`",
        "",
        f"Artifact fingerprint SHA-256: `{sources.fingerprint}`",
        "",
        "## 도면",
        "",
        "![World XY 평면도](assets/plan-view.png)",
        "",
        "![World XYZ 입체도](assets/isometric-view.png)",
        "",
        "도면은 형상과 배치 관계를 보여준다. 정확한 값은 아래 표와 JSON을 사용한다.",
        "",
        "## 좌표계",
        "",
        "World XYZ는 오른손 좌표계다. X와 Y는 바닥 평면이고 +Z는 바닥에서 상단을 향한다. "
        "거리 단위는 m이고 각도 단위는 deg다.",
        "",
        "```text",
        "direction(angle) = cos(angle) * u0 + sin(angle) * u90",
        "world_point = p0 + distance * direction(angle)",
        "rotation_axis = u0 x u90",
        "scan_plane(a, b) = p0 + a * u0 + b * u90",
        "```",
        "",
        "`p0`는 센서 측정 원점이다. `u0`는 0 deg 광선 방향이고 `u90`은 90 deg 광선 "
        "방향이다. 회전축은 오른손 기준 `u0 x u90`이다.",
        "",
        "## 적재 공간",
        "",
        "| 항목 | 값 |",
        "| --- | --- |",
        f"| Environment ID | `{environment.environment_id}` |",
        f"| X 범위 | {minimum_x:.3f} - {maximum_x:.3f} m |",
        f"| Y 범위 | {minimum_y:.3f} - {maximum_y:.3f} m |",
        f"| 바닥 Z | {environment.floor_z_m:.3f} m |",
        f"| 외벽 상단 Z | {environment.top_z_m:.3f} m |",
        f"| 유효 높이 | {height:.3f} m |",
        f"| 평면 면적 | {area:.3f} m2 |",
        f"| 전체 용량 | {area * height:.3f} m3 |",
        "",
        "경계 꼭짓점은 V1부터 순서대로 연결하고 마지막 V6을 V1에 연결한다.",
        "",
        "| 꼭짓점 | World X m | World Y m |",
        "| --- | ---: | ---: |",
    ]
    lines.extend(
        f"| V{index + 1} | {_number(x)} | {_number(y)} |" for index, (x, y) in enumerate(boundary)
    )
    lines.extend(
        [
            "",
            "## 투입구",
            "",
            "투입구 설정은 적재 표면 부피 증가의 중심인 World XY만 정의한다. Z 위치와 "
            "컨베이어 형상은 정의하지 않는다.",
            "",
            "| 투입구 | World X m | World Y m | 설정 배열 index |",
            "| --- | ---: | ---: | ---: |",
        ]
    )
    lines.extend(
        f"| INLET {index + 1} | {_number(x)} | {_number(y)} | {index} |"
        for index, (x, y) in enumerate(sources.inputs.simulator.scenario.inlet_positions_xy_m)
    )
    lines.extend(
        [
            "",
            "## 센서 설치와 바닥 측정선",
            "",
            "바닥 측정선은 센서 회전면과 `Z = floor_z_m` 평면의 교선을 적재 공간 "
            "경계로 자른 선분이다.",
            "",
            "| 센서 | 원점 p0 m | 0 deg 방향 u0 | 90 deg 방향 u90 | 회전축 u0 x u90 |",
            "| --- | --- | --- | --- | --- |",
        ]
    )
    sensors = sources.environment_document["sensors"]
    for sensor in sensors:
        lines.append(
            f"| `{sensor['sensor_id']}` | {_vector(sensor['p0_m'])} | "
            f"{_vector(sensor['u0'])} | {_vector(sensor['u90'])} | "
            f"{_vector(sensor_rotation_axis(sensor))} |"
        )
    lines.extend(
        [
            "",
            "| 센서 | 0 deg 바닥 교점 m | 바닥 교선 시작 m | 바닥 교선 끝 m |",
            "| --- | --- | --- | --- |",
        ]
    )
    for sensor in sensors:
        hit = zero_degree_floor_hit(sensor, environment.floor_z_m)
        segments = floor_measurement_segments(sensor, boundary, environment.floor_z_m)
        if len(segments) != 1:
            raise ValueError(f"expected one floor measurement segment for {sensor['sensor_id']}")
        start, end = segments[0]
        lines.append(
            f"| `{sensor['sensor_id']}` | {_vector(hit)} | {_vector(start)} | {_vector(end)} |"
        )
    lines.extend(
        [
            "",
            "## 정확한 입력 데이터",
            "",
            "### 환경과 센서 설치",
            "",
            f"Source: `{_relative_path(sources.environment_path)}`",
            "",
            "```json",
            json.dumps(sources.environment_document, ensure_ascii=True, allow_nan=False, indent=2),
            "```",
            "",
            "### 생성 시나리오와 측정 모델",
            "",
            f"Source: `{_relative_path(sources.simulator_path)}`",
            "",
            "```json",
            json.dumps(sources.simulator_document, ensure_ascii=True, allow_nan=False, indent=2),
            "```",
            "",
            "### 적재 표면 계산 정책",
            "",
            "| 항목 | 값 |",
            "| --- | ---: |",
            f"| 합성 안식각 | {DEFAULT_ANGLE_OF_REPOSE_DEG:.3f} deg |",
            f"| 경사 이완 반복 상한 | {DEFAULT_SLOPE_RELAXATION_MAX_ITERATIONS} |",
            "",
            "### 센서별 품질 분포",
            "",
            f"Source: `{_relative_path(sources.quality_path)}`",
            "",
            "```json",
            json.dumps(sources.quality_document, ensure_ascii=True, allow_nan=False, indent=2),
            "```",
            "",
        ]
    )
    output_path.write_text("\n".join(lines), encoding="utf-8")


def write_source_manifest(sources: SpecificationSources, output_path: Path) -> None:
    """Write portable provenance for the generated specification bundle."""
    document = {
        "schema_version": 2,
        "environment_id": sources.inputs.environment.environment_id,
        "input_fingerprint_sha256": sources.input_fingerprint,
        "artifact_fingerprint_sha256": sources.fingerprint,
        "sources": [
            {"path": name, "sha256": digest} for name, digest in sorted(sources.hashes.items())
        ],
        "document_generator": {
            "path": _relative_path(Path(__file__)),
            "sha256": sources.document_generator_sha256,
        },
        "model_values": sources.model_values,
        "generated_artifacts": [
            f"assets/{_PLAN_FILENAME}",
            f"assets/{_ISOMETRIC_FILENAME}",
            _MARKDOWN_FILENAME,
            f"{_DOCUMENT_BASENAME}.docx",
            f"{_DOCUMENT_BASENAME}.pdf",
        ],
    }
    output_path.write_text(
        json.dumps(document, ensure_ascii=True, allow_nan=False, indent=2) + "\n",
        encoding="utf-8",
    )


def render_pdf(docx_path: Path, output_path: Path, libreoffice: str) -> None:
    """Render the generated DOCX through LibreOffice."""
    executable = shutil.which(libreoffice)
    if executable is None:
        raise RuntimeError(f"LibreOffice executable not found: {libreoffice}")
    with tempfile.TemporaryDirectory(prefix="synthetic-environment-pdf-") as temporary:
        destination = Path(temporary)
        result = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(destination),
                str(docx_path),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        rendered = destination / f"{docx_path.stem}.pdf"
        if result.returncode != 0 or not rendered.is_file():
            detail = (result.stderr or result.stdout).strip()
            raise RuntimeError(f"LibreOffice PDF rendering failed: {detail}")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(rendered, output_path)


def check_artifacts(sources: SpecificationSources, output_dir: Path, assets_dir: Path) -> None:
    """Reject missing artifacts or artifacts derived from older public inputs."""
    paths = (
        assets_dir / _PLAN_FILENAME,
        assets_dir / _ISOMETRIC_FILENAME,
        output_dir / _MARKDOWN_FILENAME,
        output_dir / f"{_DOCUMENT_BASENAME}.docx",
        output_dir / f"{_DOCUMENT_BASENAME}.pdf",
        output_dir / _SOURCE_FILENAME,
    )
    failures = [path for path in paths if not _contains_fingerprint(path, sources.fingerprint)]
    if failures:
        joined = ", ".join(_relative_path(path) for path in failures)
        raise RuntimeError(f"missing or stale synthetic environment artifacts: {joined}")


def _draw_plan(sources: SpecificationSources, output_path: Path) -> None:
    width, height = 1800, 1450
    image = Image.new("RGB", (width * _DRAWING_SCALE, height * _DRAWING_SCALE), _WHITE)
    draw = _ScaledDraw(image)
    boundary = _boundary(sources)
    minimum_x, maximum_x, minimum_y, maximum_y = _bounds(boundary)
    left, top, right, bottom = 260, 230, 1210, 1210
    scale = min((right - left) / (maximum_x - minimum_x), (bottom - top) / (maximum_y - minimum_y))

    def point(value: Point2) -> PixelPoint:
        return (
            left + (value[0] - minimum_x) * scale,
            bottom - (value[1] - minimum_y) * scale,
        )

    _drawing_title(draw, "PLAN VIEW - WORLD XY", "unit: m | right-handed, Z up")
    _draw_plan_grid(draw, point, minimum_x, maximum_x, minimum_y, maximum_y)
    polygon = [point(value) for value in boundary]
    draw.polygon(polygon, fill=_BOUNDARY_FILL, outline=_BOUNDARY, width=6)
    _draw_boundary_vertices(draw, boundary, point)
    _draw_plan_dimensions(draw, point, minimum_x, maximum_x, minimum_y, maximum_y)
    _draw_plan_inlets(draw, sources, point)
    _draw_plan_sensors(draw, sources, point)
    _draw_plan_axes(draw, (1375, 760), 160)
    _draw_plan_legend(draw, (1290, 210))
    _draw_fingerprint(draw, sources.fingerprint, (1290, 1235))
    _save_png(image, output_path, sources.fingerprint)


def _draw_plan_grid(
    draw: _ScaledDraw,
    project: Any,
    minimum_x: float,
    maximum_x: float,
    minimum_y: float,
    maximum_y: float,
) -> None:
    grid_font = _font(24)
    for x_tenth in range(math.floor(minimum_x * 2), math.ceil(maximum_x * 2) + 1):
        x = x_tenth / 2
        start = project((x, minimum_y))
        end = project((x, maximum_y))
        draw.line((start, end), fill=_GRID, width=2)
        draw.text((start[0], start[1] + 16), _number(x), fill=_MUTED, font=grid_font, anchor="ma")
    for y_tenth in range(math.floor(minimum_y * 2), math.ceil(maximum_y * 2) + 1):
        y = y_tenth / 2
        start = project((minimum_x, y))
        end = project((maximum_x, y))
        draw.line((start, end), fill=_GRID, width=2)
        draw.text((start[0] - 18, start[1]), _number(y), fill=_MUTED, font=grid_font, anchor="rm")


def _draw_boundary_vertices(
    draw: _ScaledDraw,
    boundary: Sequence[Point2],
    project: Any,
) -> None:
    offsets = ((-12, 20), (12, 20), (12, -54), (-8, -58), (12, 20), (-12, 20))
    font = _font(24)
    for index, coordinate in enumerate(boundary):
        pixel = project(coordinate)
        draw.ellipse(_centered_box(pixel, 8), fill=_BOUNDARY, outline=_WHITE, width=2)
        offset = offsets[index] if index < len(offsets) else (12, 12)
        anchor = "ra" if offset[0] < 0 else "la"
        label = f"V{index + 1}"
        _boxed_text(
            draw,
            (pixel[0] + offset[0], pixel[1] + offset[1]),
            label,
            font=font,
            fill=_INK,
            anchor=anchor,
        )


def _draw_plan_dimensions(
    draw: _ScaledDraw,
    project: Any,
    minimum_x: float,
    maximum_x: float,
    minimum_y: float,
    maximum_y: float,
) -> None:
    horizontal_left = project((minimum_x, minimum_y))
    horizontal_right = project((maximum_x, minimum_y))
    y = horizontal_left[1] + 92
    _dimension(
        draw, (horizontal_left[0], y), (horizontal_right[0], y), f"{maximum_x - minimum_x:.3f} m"
    )
    vertical_bottom = project((minimum_x, minimum_y))
    vertical_top = project((minimum_x, maximum_y))
    x = vertical_bottom[0] - 112
    _dimension(
        draw, (x, vertical_bottom[1]), (x, vertical_top[1]), f"{maximum_y - minimum_y:.3f} m"
    )


def _draw_plan_inlets(
    draw: _ScaledDraw,
    sources: SpecificationSources,
    project: Any,
) -> None:
    font = _font(24)
    positions = sources.inputs.simulator.scenario.inlet_positions_xy_m
    for index, coordinate in enumerate(positions):
        pixel = project(coordinate)
        draw.ellipse(_centered_box(pixel, 13), fill=_INLET, outline=_WHITE, width=3)
        draw.line((pixel[0] - 19, pixel[1], pixel[0] + 19, pixel[1]), fill=_INLET, width=4)
        draw.line((pixel[0], pixel[1] - 19, pixel[0], pixel[1] + 19), fill=_INLET, width=4)
        label = f"INLET {index + 1} ({coordinate[0]:.3f}, {coordinate[1]:.3f})"
        _boxed_text(draw, (pixel[0] + 20, pixel[1] - 28), label, font=font, fill=_INLET)


def _draw_plan_sensors(
    draw: _ScaledDraw,
    sources: SpecificationSources,
    project: Any,
) -> None:
    font = _font(23)
    sensors = sources.environment_document["sensors"]
    for index, sensor in enumerate(sensors):
        p0 = tuple(float(value) for value in sensor["p0_m"])
        sensor_pixel = project((p0[0], p0[1]))
        _diamond(draw, sensor_pixel, radius=14, fill=_SENSOR)
        label_offset = (22, 30) if index == 0 else (20, 52)
        anchor = "la"
        _boxed_text(
            draw,
            (sensor_pixel[0] + label_offset[0], sensor_pixel[1] + label_offset[1]),
            f"SENSOR {index + 1}",
            font=font,
            fill=_SENSOR,
            anchor=anchor,
        )


def _draw_plan_axes(draw: _ScaledDraw, origin: PixelPoint, scale: float) -> None:
    _arrow(draw, origin, (origin[0] + scale * 0.8, origin[1]), fill="#b91c1c", width=5)
    _arrow(draw, origin, (origin[0], origin[1] - scale * 0.8), fill="#15803d", width=5)
    draw.text(
        (origin[0] + scale * 0.85, origin[1]), "+X", fill="#b91c1c", font=_font(26), anchor="lm"
    )
    draw.text(
        (origin[0], origin[1] - scale * 0.85), "+Y", fill="#15803d", font=_font(26), anchor="mb"
    )


def _draw_plan_legend(draw: _ScaledDraw, origin: PixelPoint) -> None:
    x, y = origin
    draw.rounded_rectangle(
        (x, y, x + 430, y + 270), radius=18, fill="#f8fafc", outline=_GRID, width=3
    )
    draw.text((x + 24, y + 22), "LEGEND", fill=_INK, font=_font(30))
    rows = (
        (_BOUNDARY, "Storage boundary"),
        (_SENSOR, "Sensor XY projection"),
        (_INLET, "Inlet XY position"),
    )
    for index, (color, label) in enumerate(rows):
        row_y = y + 88 + index * 57
        draw.line((x + 30, row_y, x + 88, row_y), fill=color, width=8)
        draw.text((x + 112, row_y), label, fill=_INK, font=_font(24), anchor="lm")
    draw.text(
        (x + 24, y + 238),
        "Exact coordinates: Markdown specification",
        fill=_MUTED,
        font=_font(16),
    )


def _draw_isometric(sources: SpecificationSources, output_path: Path) -> None:
    width, height = 1800, 1600
    image = Image.new("RGB", (width * _DRAWING_SCALE, height * _DRAWING_SCALE), _WHITE)
    draw = _ScaledDraw(image)
    boundary = _boundary(sources)
    floor_z = sources.inputs.environment.floor_z_m
    top_z = sources.inputs.environment.top_z_m
    raw_points = [_iso_raw((x, y, z)) for x, y in boundary for z in (floor_z, top_z)]
    minimum_u = min(point[0] for point in raw_points)
    maximum_u = max(point[0] for point in raw_points)
    minimum_v = min(point[1] for point in raw_points)
    maximum_v = max(point[1] for point in raw_points)
    left, top, right, bottom = 100, 160, 1190, 1450
    scale = min(
        (right - left) / (maximum_u - minimum_u),
        (bottom - top) / (maximum_v - minimum_v),
    )
    drawing_width = (maximum_u - minimum_u) * scale
    drawing_height = (maximum_v - minimum_v) * scale
    offset_x = left + (right - left - drawing_width) / 2
    offset_y = top + (bottom - top - drawing_height) / 2

    def project(value: Point3) -> PixelPoint:
        u, v = _iso_raw(value)
        return (
            offset_x + (u - minimum_u) * scale,
            offset_y + (v - minimum_v) * scale,
        )

    _drawing_title(
        draw,
        "ISOMETRIC VIEW - WORLD XYZ",
        "orthographic technical projection | unit: m | empty-floor reference",
    )
    floor = [project((x, y, floor_z)) for x, y in boundary]
    ceiling = [project((x, y, top_z)) for x, y in boundary]
    draw.polygon(floor, fill="#eef3f7")
    _draw_isometric_walls(image, boundary, floor_z, top_z, project)
    draw = _ScaledDraw(image)
    draw.line((*floor, floor[0]), fill="#34495e", width=6, joint="curve")
    draw.line((*ceiling, ceiling[0]), fill="#34495e", width=5, joint="curve")
    _draw_isometric_inlets(draw, sources, project)
    _draw_isometric_sensor_planes(image, sources, project)
    draw = _ScaledDraw(image)
    _draw_isometric_axes(draw, project)
    _draw_isometric_dimensions(draw, boundary, floor_z, top_z, project)
    _draw_isometric_legend(draw, (1270, 190))
    _draw_fingerprint(draw, sources.fingerprint, (1260, 1450))
    _save_png(image, output_path, sources.fingerprint)


def _draw_isometric_walls(
    image: Image.Image,
    boundary: Sequence[Point2],
    floor_z: float,
    top_z: float,
    project: Any,
) -> None:
    for first, second in _polygon_edges(boundary):
        wall = (
            project((first[0], first[1], floor_z)),
            project((second[0], second[1], floor_z)),
            project((second[0], second[1], top_z)),
            project((first[0], first[1], top_z)),
        )
        _translucent_polygon(image, wall, (183, 207, 224, 38))
    draw = _ScaledDraw(image)
    for x, y in boundary:
        draw.line(
            (project((x, y, floor_z)), project((x, y, top_z))),
            fill="#34495e",
            width=5,
        )


def _draw_isometric_inlets(
    draw: _ScaledDraw,
    sources: SpecificationSources,
    project: Any,
) -> None:
    floor_z = sources.inputs.environment.floor_z_m
    top_z = sources.inputs.environment.top_z_m
    for index, (x, y) in enumerate(sources.inputs.simulator.scenario.inlet_positions_xy_m):
        low = project((x, y, floor_z))
        high = project((x, y, top_z))
        _dashed_line(draw, low, high, fill=_INLET, width=4, dash=14)
        draw.ellipse(_centered_box(high, 10), fill=_INLET, outline=_WHITE, width=3)
        _boxed_text(
            draw,
            (high[0] + 14, high[1] + 24),
            f"INLET {index + 1}",
            font=_font(19),
            fill=_INLET,
        )


def _draw_isometric_sensor_planes(
    image: Image.Image,
    sources: SpecificationSources,
    project: Any,
) -> None:
    floor_z = sources.inputs.environment.floor_z_m
    boundary = _boundary(sources)
    colors = (_SENSOR_ONE, _SENSOR_TWO)
    fills = ((25, 118, 210, 28), (0, 150, 136, 28))
    for index, sensor in enumerate(sources.environment_document["sensors"]):
        p0 = tuple(float(value) for value in sensor["p0_m"])
        hit = zero_degree_floor_hit(sensor, floor_z)
        segments = floor_measurement_segments(sensor, boundary, floor_z)
        if len(segments) != 1:
            raise ValueError(f"expected one floor measurement segment for {sensor['sensor_id']}")
        start_3d, end_3d = segments[0]
        source = project(p0)
        start = project(start_3d)
        end = project(end_3d)
        floor_hit = project(hit)
        color = colors[index]
        _translucent_polygon(image, (source, start, end), fills[index])
        draw = _ScaledDraw(image)
        _dashed_line(draw, source, start, fill=color, width=3, dash=15)
        _dashed_line(draw, source, end, fill=color, width=3, dash=15)
        _dashed_line(draw, source, floor_hit, fill=color, width=4, dash=11)
        draw.line((start, end), fill=color, width=9)
        draw.ellipse(_centered_box(start, 7), fill=_WHITE, outline=color, width=4)
        draw.ellipse(_centered_box(end, 7), fill=_WHITE, outline=color, width=4)
        draw.ellipse(_centered_box(floor_hit, 9), fill=_WHITE, outline=color, width=4)
        draw.ellipse(_centered_box(source, 15), fill="#34495e", outline=_WHITE, width=3)
        draw.ellipse(_centered_box(source, 7), fill=color)
        _boxed_text(
            draw,
            (source[0] + (-22 if index == 0 else 22), source[1] - 30),
            f"SENSOR {index + 1}",
            font=_font(23),
            fill=color,
            anchor="ra" if index == 0 else "la",
        )


def _draw_isometric_axes(draw: _ScaledDraw, project: Any) -> None:
    origin = project((0.0, 0.0, 0.0))
    axes = (
        ((0.8, 0.0, 0.0), "X", "#b91c1c"),
        ((0.0, 0.8, 0.0), "Y", "#15803d"),
        ((0.0, 0.0, 0.8), "Z", "#1d4ed8"),
    )
    for endpoint, label, color in axes:
        tip = project(endpoint)
        _arrow(draw, origin, tip, fill=color, width=5)
        draw.text(tip, f"+{label}", fill=color, font=_font(22), anchor="mm")


def _draw_isometric_dimensions(
    draw: _ScaledDraw,
    boundary: Sequence[Point2],
    floor_z: float,
    top_z: float,
    project: Any,
) -> None:
    minimum_x, maximum_x, minimum_y, maximum_y = _bounds(boundary)
    x_start = project((minimum_x, minimum_y, floor_z))
    x_end = project((maximum_x, minimum_y, floor_z))
    _dimension(
        draw,
        (x_start[0], x_start[1] + 72),
        (x_end[0], x_end[1] + 72),
        f"X {maximum_x - minimum_x:.3f} m",
    )
    y_start = project((maximum_x, minimum_y, floor_z))
    y_end = project((maximum_x, maximum_y, floor_z))
    _dimension(
        draw,
        (y_start[0] + 66, y_start[1] + 18),
        (y_end[0] + 66, y_end[1] + 18),
        f"Y {maximum_y - minimum_y:.3f} m",
    )
    z_start = project((maximum_x, maximum_y, floor_z))
    z_end = project((maximum_x, maximum_y, top_z))
    _dimension(
        draw,
        (z_start[0] - 56, z_start[1]),
        (z_end[0] - 56, z_end[1]),
        f"Z {top_z - floor_z:.3f} m",
    )


def _draw_isometric_legend(
    draw: _ScaledDraw,
    origin: PixelPoint,
) -> None:
    x, y = origin
    right = x + 455
    bottom = y + 365
    draw.rounded_rectangle((x, y, right, bottom), radius=18, fill="#f8fafc", outline=_GRID, width=3)
    draw.text((x + 24, y + 22), "DRAWING KEY", fill=_INK, font=_font(28))
    rows = (
        ("#34495e", "Storage wall and floor"),
        (_INLET, "Inlet XY axis"),
        (_SENSOR_ONE, "SENSOR 1 scan plane and floor line"),
        (_SENSOR_TWO, "SENSOR 2 scan plane and floor line"),
    )
    for index, (color, label) in enumerate(rows):
        row_y = y + 82 + index * 44
        draw.line((x + 28, row_y, x + 76, row_y), fill=color, width=7)
        draw.text((x + 94, row_y), label, fill=_INK, font=_font(19), anchor="lm")

    draw.text(
        (x + 24, bottom - 48),
        "Exact coordinates: Markdown specification",
        fill=_MUTED,
        font=_font(16),
    )


def _translucent_polygon(
    image: Image.Image,
    points: Sequence[PixelPoint],
    fill: tuple[int, int, int, int],
) -> None:
    overlay = Image.new("RGBA", image.size, (255, 255, 255, 0))
    ImageDraw.Draw(overlay).polygon(_scale_coordinates(points), fill=fill)
    image.paste(overlay, (0, 0), overlay)


def _write_title_page(document: DocumentType, sources: SpecificationSources) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(4.2)
    run = paragraph.add_run("공개 합성 LiDAR 시뮬레이션 환경 규격서")
    run.font.size = Pt(26)
    run.font.bold = False
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(str(sources.inputs.environment.environment_id)).font.size = Pt(16)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(1.2)
    paragraph.add_run("적재 공간, 센서 설치, 투입구, 시뮬레이션 및 공간 좌표 규격").font.size = Pt(
        13
    )
    _add_notice(
        document,
        "이 규격은 실제 현장 정보가 아닌 공개 가능한 합성 개발 환경을 나타낸다. "
        "수치의 정본은 versioned JSON 3개이며 이 문서와 도면은 해당 입력에서 생성한 파생 산출물이다.",
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(2.0)
    paragraph.add_run(f"Artifact fingerprint SHA-256\n{sources.fingerprint}").font.size = Pt(9)
    _page_break(document)


def _write_source_section(document: DocumentType, sources: SpecificationSources) -> None:
    _heading(document, "1. 정본과 적용 범위", 1)
    _paragraph(
        document,
        "문서가 나타내는 구성 요소는 적재 공간, 투입구, LiDAR 센서 2대, 적재 표면 모델과 "
        "센서 측정 모델의 5개다. 문서는 합성 통합 검증에서 사용하는 현재 설정만 기술한다.",
    )
    _add_table(
        document,
        ("정본", "책임", "SHA-256"),
        tuple(
            (
                name,
                _source_responsibility(name),
                digest,
            )
            for name, digest in sources.hashes.items()
        ),
    )
    _paragraph(
        document,
        "정본과 문서가 다르면 JSON을 우선한다. JSON이 바뀌면 생성 도구가 fingerprint 불일치를 "
        "검출하므로 도면, DOCX와 PDF를 함께 다시 생성해야 한다.",
    )


def _write_coordinate_section(document: DocumentType, sources: SpecificationSources) -> None:
    _heading(document, "2. 좌표계와 단위", 1)
    _paragraph(
        document,
        "좌표 표현은 world와 sensor scan의 2개다. World가 적재 공간, 투입구와 sensor 설치 "
        "형상의 공통 정본이고 sensor scan은 광선의 회전 각도와 거리를 나타낸다.",
    )
    _add_table(
        document,
        ("표현", "단위", "축과 의미", "사용 경계"),
        (
            ("World XYZ", "m, deg", "오른손 좌표, +Z 상단", "환경, 적재 표면, 관찰 stream"),
            ("Sensor polar", "mm, mdeg", "sensor 원점의 회전 각도와 거리", "gRPC ScanFrame"),
        ),
    )
    _paragraph(document, "Sensor 광선 방향은 다음 식으로 계산한다.")
    _formula(document, "direction(angle) = cos(angle) * u0 + sin(angle) * u90")
    _formula(document, "world_point = p0 + distance * direction(angle)")
    _paragraph(
        document,
        "Sensor 회전 형상은 scan plane, 회전축, 0 deg 바닥 교점과 바닥 측정선의 4개 파생값으로 "
        "표현한다.",
    )
    _formula(document, "scan_plane(a, b) = p0 + a * u0 + b * u90")
    _formula(document, "rotation_axis = u0 x u90")
    _formula(document, "H0 = p0 + ((floor_z - p0.z) / u0.z) * u0")
    _formula(document, "floor_measurement_line(t) = H0 + t * u90")
    _paragraph(
        document,
        "World 원점은 경계 꼭짓점 V1의 바닥 위치 (0, 0, 0)이다. XY는 평면이고 Z는 절대 "
        "높이다. 각도는 degree이며 angle 0과 angle 90의 방향은 각각 u0과 u90이다.",
    )


def _write_space_section(
    document: DocumentType,
    sources: SpecificationSources,
    drawings: DrawingPaths,
) -> None:
    _heading(document, "3. 적재 공간", 1)
    environment = sources.inputs.environment
    boundary = tuple(environment.boundary_xy_m)
    minimum_x, maximum_x, minimum_y, maximum_y = _bounds(boundary)
    area = polygon_area(boundary)
    height = environment.top_z_m - environment.floor_z_m
    _add_table(
        document,
        ("항목", "값", "의미"),
        (
            ("Environment ID", environment.environment_id, "합성 환경 식별자"),
            ("X 범위", f"{minimum_x:.3f} - {maximum_x:.3f} m", "평면 bounding 범위"),
            ("Y 범위", f"{minimum_y:.3f} - {maximum_y:.3f} m", "평면 bounding 범위"),
            ("바닥 Z", f"{environment.floor_z_m:.3f} m", "절대 바닥 높이"),
            ("외벽 상단 Z", f"{environment.top_z_m:.3f} m", "절대 높이 상한"),
            ("유효 높이", f"{height:.3f} m", "상단 Z - 바닥 Z"),
            ("평면 면적", f"{area:.3f} m2", "경계 polygon 면적"),
            ("전체 용량", f"{area * height:.3f} m3", "평면 면적 x 유효 높이"),
        ),
    )
    rows = []
    for index, (first, second) in enumerate(
        zip(boundary, boundary[1:] + boundary[:1], strict=True)
    ):
        next_index = (index + 1) % len(boundary) + 1
        rows.append(
            (
                f"V{index + 1}",
                f"({_number(first[0])}, {_number(first[1])})",
                f"V{index + 1} -> V{next_index}",
                f"{math.dist(first, second):.6f} m",
            )
        )
    _add_table(document, ("꼭짓점", "World XY m", "다음 변", "변 길이"), tuple(rows))
    _figure(document, drawings.plan, "그림 1. World XY 경계 형상, 주요 치수와 배치 관계.")


def _write_sensor_section(
    document: DocumentType,
    sources: SpecificationSources,
    drawings: DrawingPaths,
) -> None:
    _heading(document, "4. 투입구와 센서 설치", 1)
    inlets = sources.inputs.simulator.scenario.inlet_positions_xy_m
    _paragraph(
        document,
        "투입구 설정은 XY 위치만 정의한다. 투입 높이와 설비 형상은 정의하지 않으며 시뮬레이션은 "
        "해당 XY를 표면 부피 증가의 중심으로 사용한다.",
    )
    _add_table(
        document,
        ("투입구", "World X m", "World Y m", "배열 index"),
        tuple(
            (f"INLET {index + 1}", _number(x), _number(y), str(index))
            for index, (x, y) in enumerate(inlets)
        ),
    )
    sensor_rows = []
    measurement_rows = []
    boundary = _boundary(sources)
    for sensor in sources.environment_document["sensors"]:
        hit = zero_degree_floor_hit(sensor, sources.inputs.environment.floor_z_m)
        segments = floor_measurement_segments(
            sensor,
            boundary,
            sources.inputs.environment.floor_z_m,
        )
        if len(segments) != 1:
            raise ValueError(f"expected one floor measurement segment for {sensor['sensor_id']}")
        segment_start, segment_end = segments[0]
        sensor_rows.append(
            (
                sensor["sensor_id"],
                _vector(sensor["p0_m"]),
                _vector(sensor["u0"]),
                _vector(sensor["u90"]),
                f"{sensor_tilt_from_down_deg(sensor):.6f}",
                _vector(hit),
            )
        )
        measurement_rows.append(
            (
                sensor["sensor_id"],
                _vector(sensor_rotation_axis(sensor)),
                f"H0 + t * {_vector(sensor['u90'])}",
                _vector(segment_start),
                _vector(segment_end),
            )
        )
    _add_table(
        document,
        ("Sensor", "p0 m", "u0", "u90", "수직 하향 대비 deg", "0 deg 바닥 교점 m"),
        tuple(sensor_rows),
    )
    _add_table(
        document,
        ("Sensor", "회전축 u0 x u90", "바닥 측정선", "경계 교점 A m", "경계 교점 B m"),
        tuple(measurement_rows),
    )
    _paragraph(
        document,
        "p0는 sensor 원점이고 u0와 u90은 scan plane의 직교 단위벡터다. 회전축은 오른손 기준 "
        "u0 x u90이다. 바닥 측정선은 360 deg scan plane과 floor_z_m 평면의 교선을 적재 공간 "
        "경계로 자른 선분이다. 입체도는 외벽, sensor 회전면과 빈 바닥 교선을 표시한다. "
        "방향벡터와 교점은 위 표의 좌표를 기준으로 한다.",
    )
    _figure(
        document,
        drawings.isometric,
        "그림 2. 외벽, sensor 회전면과 빈 바닥 측정선의 입체도.",
    )


def _write_simulation_section(document: DocumentType, sources: SpecificationSources) -> None:
    _heading(document, "5. 적재 표면과 시나리오", 1)
    scenario = sources.inputs.simulator.scenario
    surface = scenario.surface
    boundary = tuple(sources.inputs.environment.boundary_xy_m)
    minimum_x, maximum_x, minimum_y, maximum_y = _bounds(boundary)
    x_nodes = math.ceil((maximum_x - minimum_x) / surface.cell_size_m) + 1
    y_nodes = math.ceil((maximum_y - minimum_y) / surface.cell_size_m) + 1
    _add_table(
        document,
        ("항목", "값", "적용 의미"),
        (
            ("평균 적재 시간", f"{scenario.mean_fill_duration_s:.3f} s", "24시간 기준"),
            (
                "회차별 적재 시간 계수",
                _range(scenario.fill_duration_factor_range),
                "평균 대비 범위",
            ),
            ("적재 속도 계수", _range(scenario.fill_rate_factor_range), "구간별 속도 범위"),
            (
                "적재 속도 구간",
                _range(scenario.fill_rate_change_duration_s_range, "s"),
                "구간 지속 시간",
            ),
            ("수거 시작 적재율", _range(scenario.collection_threshold_range), "회차별 임계치"),
            (
                "수거 시간 계수",
                _range(scenario.collection_duration_factor_range),
                "평균 적재 시간 대비",
            ),
            ("수거 속도 계수", _range(scenario.collection_rate_factor_range), "구간별 속도 범위"),
            (
                "수거 속도 구간",
                _range(scenario.collection_rate_change_duration_s_range, "s"),
                "구간 지속 시간",
            ),
            ("표면 cell", f"{surface.cell_size_m:.3f} m", "정규 격자 간격"),
            ("표면 node shape", f"Y {y_nodes} x X {x_nodes}", "배열 순서와 node 수"),
            ("표면 갱신", f"{surface.update_interval_s:.3f} s", "시뮬레이션 시각 기준"),
            ("투입 확산 반경", f"{surface.pile_spread_radius_m:.3f} m", "국소 부피 분포"),
            ("요철 높이", _range(surface.roughness_height_range_m, "m"), "signed peak 변화"),
            ("요철 반경", _range(surface.roughness_radius_range_m, "m"), "국소 변화 범위"),
            ("합성 안식각", f"{DEFAULT_ANGLE_OF_REPOSE_DEG:.3f} deg", "경사 이완 상한"),
            (
                "경사 이완 반복 상한",
                str(DEFAULT_SLOPE_RELAXATION_MAX_ITERATIONS),
                "표면 갱신당 반복 수",
            ),
            (
                "투입구 전환 활성 비율",
                f"{scenario.inlet_switch_activation_ratio:.3f}",
                "회차 진행 비율",
            ),
            (
                "투입구 전환 높이 차",
                f"{scenario.inlet_switch_height_difference_m:.3f} m",
                "전환 판정",
            ),
            ("투입구 비교 반경", f"{scenario.inlet_comparison_radius_m:.3f} m", "평균 높이 범위"),
        ),
    )
    _paragraph(
        document,
        "높이 배열은 Y-major 순서이며 polygon의 bounding box를 덮는다. 마지막 Y 좌표는 "
        f"{minimum_y + (y_nodes - 1) * surface.cell_size_m:.3f} m이고 polygon 밖의 node는 "
        "렌더링과 부피 계산에서 경계로 잘라낸다.",
    )


def _write_measurement_section(document: DocumentType, sources: SpecificationSources) -> None:
    _heading(document, "6. LiDAR 측정 모델", 1)
    measurement = sources.inputs.simulator.measurement
    nominal_points = measurement.sample_rate_hz / measurement.rotation_rate_hz
    _add_table(
        document,
        ("항목", "Sensor 1대", "Sensor 2대 합계"),
        (
            (
                "Sample rate",
                f"{measurement.sample_rate_hz:.3f} point/s",
                f"{2 * measurement.sample_rate_hz:.3f} point/s",
            ),
            (
                "Rotation rate",
                f"{measurement.rotation_rate_hz:.3f} scan/s",
                f"{2 * measurement.rotation_rate_hz:.3f} scan/s",
            ),
            ("명목 회전당 측정점", f"{nominal_points:.3f}", "고정 wire 계약 아님"),
            (
                "측정 거리",
                f"{measurement.min_distance_m:.3f} - {measurement.max_distance_m:.3f} m",
                "동일",
            ),
            (
                "거리 noise 표준편차",
                f"{measurement.distance_noise.standard_deviation_m:.3f} m",
                "독립 적용",
            ),
            ("거리 noise 절대 상한", f"{measurement.distance_noise.limit_m:.3f} m", "독립 적용"),
        ),
    )
    distortions = measurement.distortions
    _add_table(
        document,
        ("합성 현상", "활성", "빈도 또는 비율", "공간 또는 거리", "지속 시간"),
        (
            (
                "낙하물",
                _enabled(distortions.falling_material.enabled),
                f"{distortions.falling_material.event_rate_per_s:.3f} event/s",
                f"반경 {_range(distortions.falling_material.radius_m_range, 'm')}, 거리 감소 {_range(distortions.falling_material.distance_reduction_m_range, 'm')}",
                _range(distortions.falling_material.duration_s_range, "s"),
            ),
            (
                "빈틈",
                _enabled(distortions.voids.enabled),
                f"표면 {distortions.voids.surface_area_ratio:.3f}",
                f"반경 {_range(distortions.voids.radius_m_range, 'm')}, 거리 증가 {_range(distortions.voids.distance_increase_m_range, 'm')}, 덮임 높이 {distortions.voids.cover_height_increase_m:.3f} m",
                _range(distortions.voids.duration_s_range, "s"),
            ),
            (
                "수거 가림",
                _enabled(distortions.collection_occlusion.enabled),
                f"간격 {_range(distortions.collection_occlusion.event_interval_s_range, 's')}",
                f"반경 {_range(distortions.collection_occlusion.radius_m_range, 'm')}, 거리 감소 {_range(distortions.collection_occlusion.distance_reduction_m_range, 'm')}",
                _range(distortions.collection_occlusion.duration_s_range, "s"),
            ),
            (
                "반사 경로 오류",
                _enabled(distortions.reflection_error.enabled),
                f"확률 {distortions.reflection_error.probability:.6f}",
                f"거리 감소 {_range(distortions.reflection_error.distance_reduction_m_range, 'm')}",
                "point 단위",
            ),
            (
                "Sensor dropout",
                _enabled(distortions.dropout.enabled),
                f"간격 {_range(distortions.dropout.event_interval_s_range, 's')}",
                "전체 sensor",
                _range(distortions.dropout.duration_s_range, "s"),
            ),
        ),
    )
    quality_rows = []
    for sensor in sources.quality_document["sensors"]:
        quality_rows.append(
            (
                sensor["sensor_id"],
                _quality_frequencies(sensor["valid_distance_frequencies"]),
                _quality_frequencies(sensor["invalid_distance_frequencies"]),
                "HQ quality >> 2",
            )
        )
    _add_table(
        document,
        ("Sensor", "유효 거리 HQ 빈도", "무효 거리 HQ 빈도", "Wire 변환"),
        tuple(quality_rows),
    )


def _write_interpretation_section(document: DocumentType, sources: SpecificationSources) -> None:
    _heading(document, "7. 환경 해석 규칙", 1)
    _add_table(
        document,
        ("항목", "규칙"),
        (
            ("공간", "World XYZ, meter, degree를 유일한 공간 정본으로 사용"),
            ("적재 공간", "경계 polygon, floor Z와 top Z로 정의"),
            ("Sensor 설치", "p0, u0와 u90으로 위치와 scan plane을 정의"),
            ("투입구", "World XY 위치만 정의하고 Z와 설비 형상은 정의하지 않음"),
            ("적재 표면", "World XY 격자와 절대 Z 높이로 정의"),
        ),
    )
    _paragraph(
        document,
        f"이 문서의 artifact fingerprint는 {sources.fingerprint}이다.",
    )


def _configure_document(document: DocumentType, sources: SpecificationSources) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.start_type = WD_SECTION.NEW_PAGE
    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = styles[style_name]
        style.font.name = "Noto Sans CJK KR"
        style.font.bold = False
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    styles["Normal"].font.size = Pt(9.5)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 2"].font.size = Pt(13)
    core = document.core_properties
    core.title = "공개 합성 LiDAR 시뮬레이션 환경 규격서"
    core.subject = f"artifact-fingerprint:{sources.fingerprint}"
    core.author = "ajin-scrap-monitoring"
    core.keywords = f"synthetic environment,LiDAR,{sources.fingerprint}"
    core.comments = "Derived from versioned public JSON sources."
    core.created = _FIXED_DOCUMENT_TIME
    core.modified = _FIXED_DOCUMENT_TIME
    core.last_printed = _FIXED_DOCUMENT_TIME


def _add_table(
    document: DocumentType,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, value)
        _shade_cell(cell, "DCE6F1")
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            _set_cell_text(cell, str(value))
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)


def _set_cell_text(cell: Any, value: str) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value)
    run.font.name = "Noto Sans CJK KR"
    run.font.size = Pt(7.6)
    run.font.bold = False
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")


def _shade_cell(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _heading(document: DocumentType, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _paragraph(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)


def _formula(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "DejaVu Sans Mono"
    run.font.size = Pt(8.5)


def _add_notice(document: DocumentType, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    _set_cell_text(table.cell(0, 0), text)
    _shade_cell(table.cell(0, 0), "EEF4FA")


def _figure(document: DocumentType, path: Path, caption: str) -> None:
    _page_break(document)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.75))
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_break(document)


def _page_break(document: DocumentType) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _drawing_title(draw: _ScaledDraw, title: str, subtitle: str) -> None:
    draw.text((90, 52), title, fill=_INK, font=_font(42))
    draw.text((92, 105), subtitle, fill=_MUTED, font=_font(24))


def _draw_fingerprint(draw: _ScaledDraw, fingerprint: str, origin: PixelPoint) -> None:
    draw.text(
        origin,
        f"Artifact SHA-256\n{fingerprint[:32]}\n{fingerprint[32:]}",
        fill=_MUTED,
        font=_font(17),
    )


def _dimension(
    draw: _ScaledDraw,
    start: PixelPoint,
    end: PixelPoint,
    label: str,
) -> None:
    draw.line((start, end), fill=_INK, width=3)
    dx, dy = end[0] - start[0], end[1] - start[1]
    length = math.hypot(dx, dy)
    nx, ny = (-dy / length * 9, dx / length * 9)
    for point in (start, end):
        draw.line((point[0] - nx, point[1] - ny, point[0] + nx, point[1] + ny), fill=_INK, width=3)
    middle = ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2)
    _boxed_text(draw, middle, label, font=_font(24), fill=_INK, anchor="mm")


def _arrow(
    draw: _ScaledDraw,
    start: PixelPoint,
    end: PixelPoint,
    *,
    fill: str,
    width: int,
) -> None:
    draw.line((start, end), fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    head = 17 + width
    spread = math.pi / 7
    points = (
        end,
        (end[0] - head * math.cos(angle - spread), end[1] - head * math.sin(angle - spread)),
        (end[0] - head * math.cos(angle + spread), end[1] - head * math.sin(angle + spread)),
    )
    draw.polygon(points, fill=fill)


def _dashed_line(
    draw: _ScaledDraw,
    start: PixelPoint,
    end: PixelPoint,
    *,
    fill: str,
    width: int,
    dash: int,
) -> None:
    length = math.dist(start, end)
    if length == 0:
        return
    dx = (end[0] - start[0]) / length
    dy = (end[1] - start[1]) / length
    offset = 0.0
    while offset < length:
        segment_end = min(length, offset + dash)
        draw.line(
            (
                (start[0] + dx * offset, start[1] + dy * offset),
                (start[0] + dx * segment_end, start[1] + dy * segment_end),
            ),
            fill=fill,
            width=width,
        )
        offset += dash * 1.75


def _diamond(draw: _ScaledDraw, center: PixelPoint, *, radius: int, fill: str) -> None:
    x, y = center
    draw.polygon(((x, y - radius), (x + radius, y), (x, y + radius), (x - radius, y)), fill=fill)
    draw.line(
        ((x, y - radius), (x + radius, y), (x, y + radius), (x - radius, y), (x, y - radius)),
        fill=_WHITE,
        width=3,
    )


def _boxed_text(
    draw: _ScaledDraw,
    position: PixelPoint,
    text: str,
    *,
    font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
    fill: str,
    anchor: str = "la",
) -> None:
    box = draw.multiline_textbbox(position, text, font=font, anchor=anchor, spacing=3)
    padded = (box[0] - 6, box[1] - 4, box[2] + 6, box[3] + 4)
    draw.rounded_rectangle(padded, radius=5, fill=_WHITE, outline="#e2e8f0", width=1)
    draw.multiline_text(position, text, fill=fill, font=font, anchor=anchor, spacing=3)


def _centered_box(center: PixelPoint, radius: float) -> tuple[float, float, float, float]:
    return (center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius)


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    if _FONT_PATH.is_file():
        return ImageFont.truetype(str(_FONT_PATH), size=size * _DRAWING_SCALE)
    return ImageFont.load_default(size=size * _DRAWING_SCALE)


def _iso_raw(point: Point3) -> PixelPoint:
    x, y, z = point
    return (
        (x + y) * math.cos(math.pi / 6),
        (x - y) * 0.4 - z * 0.56,
    )


def _save_png(image: Image.Image, path: Path, fingerprint: str) -> None:
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Artifact-Fingerprint", fingerprint)
    metadata.add_text(
        "Source-Files",
        "examples/environment.v1.json;examples/simulator.v2.json;examples/quality-profile.v1.json",
    )
    image.save(path, format="PNG", optimize=True, pnginfo=metadata, dpi=(300, 300))


def _normalize_docx(source: Path, destination: Path) -> None:
    with (
        zipfile.ZipFile(source) as archive,
        zipfile.ZipFile(
            destination,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as normalized,
    ):
        for source_info in sorted(archive.infolist(), key=lambda info: info.filename):
            target_info = zipfile.ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            target_info.compress_type = zipfile.ZIP_DEFLATED
            target_info.external_attr = source_info.external_attr
            target_info.create_system = source_info.create_system
            normalized.writestr(target_info, archive.read(source_info.filename))


def _contains_fingerprint(path: Path, fingerprint: str) -> bool:
    if not path.is_file():
        return False
    if path.suffix == ".png":
        with Image.open(path) as image:
            return image.info.get("Artifact-Fingerprint") == fingerprint
    if path.suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            return fingerprint.encode() in archive.read("docProps/core.xml")
    if path.suffix == ".pdf":
        payload = path.read_bytes()
        markers = (
            fingerprint.encode(),
            fingerprint.encode("utf-16-be").hex().encode(),
            fingerprint.encode("utf-16-be").hex().upper().encode(),
        )
        return any(marker in payload for marker in markers)
    if path.suffix == ".json":
        return _load_json(path).get("artifact_fingerprint_sha256") == fingerprint
    if path.suffix == ".md":
        return fingerprint.encode() in path.read_bytes()
    return False


def _load_json(path: Path) -> JsonObject:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"expected JSON object: {path}")
    return value


def _relative_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(_ROOT).as_posix()
    except ValueError as error:
        raise ValueError("specification inputs must be inside the repository") from error


def _source_responsibility(name: str) -> str:
    if name.endswith("environment.v1.json"):
        return "적재 공간과 sensor 설치"
    if name.endswith("simulator.v2.json"):
        return "시나리오, 측정과 seed"
    return "Sensor별 합성 quality 분포"


def _boundary(sources: SpecificationSources) -> tuple[Point2, ...]:
    result: list[Point2] = []
    for raw_point in sources.environment_document["boundary_xy_m"]:
        if not isinstance(raw_point, list) or len(raw_point) != 2:
            raise ValueError("boundary point must contain exactly two numbers")
        result.append((float(raw_point[0]), float(raw_point[1])))
    return tuple(result)


def _bounds(boundary: Sequence[Point2]) -> tuple[float, float, float, float]:
    x_values = tuple(point[0] for point in boundary)
    y_values = tuple(point[1] for point in boundary)
    return min(x_values), max(x_values), min(y_values), max(y_values)


def _number(value: float) -> str:
    if math.isclose(float(value), round(float(value)), abs_tol=1e-12):
        return str(round(float(value)))
    return f"{float(value):.6f}".rstrip("0").rstrip(".")


def _vector(values: Iterable[float]) -> str:
    return "[" + ", ".join(f"{_clean_zero(float(value)):.15g}" for value in values) + "]"


def _fixed_vector(values: Iterable[float]) -> str:
    return "(" + ", ".join(f"{_clean_zero(float(value)):.6f}" for value in values) + ")"


def _clean_zero(value: float) -> float:
    return 0.0 if abs(value) < 0.5e-6 else value


def _range(values: Sequence[float], unit: str = "") -> str:
    suffix = f" {unit}" if unit else ""
    return f"{float(values[0]):.6g} - {float(values[1]):.6g}{suffix}"


def _enabled(value: bool) -> str:
    return "enabled" if value else "disabled"


def _quality_frequencies(value: JsonObject) -> str:
    return ", ".join(f"{quality}:{frequency}" for quality, frequency in value.items())


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--simulator-config", type=Path, default=_DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path, default=_DEFAULT_OUTPUT)
    parser.add_argument("--assets-dir", type=Path, default=_DEFAULT_ASSETS)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--render-pdf", action="store_true")
    parser.add_argument("--libreoffice", default="libreoffice")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Generate or validate the public synthetic environment specification."""
    arguments = _build_parser().parse_args(argv)
    sources = load_sources(arguments.simulator_config)
    output_dir = arguments.output_dir.resolve()
    assets_dir = arguments.assets_dir.resolve()
    if arguments.check:
        check_artifacts(sources, output_dir, assets_dir)
        print(f"synthetic_environment_spec=ok artifact_fingerprint={sources.fingerprint}")
        return 0
    drawings = generate_drawings(sources, assets_dir)
    markdown_path = output_dir / _MARKDOWN_FILENAME
    docx_path = output_dir / f"{_DOCUMENT_BASENAME}.docx"
    pdf_path = output_dir / f"{_DOCUMENT_BASENAME}.pdf"
    write_source_manifest(sources, output_dir / _SOURCE_FILENAME)
    generate_markdown(sources, markdown_path)
    generate_docx(sources, drawings, docx_path)
    print(f"synthetic_environment_markdown={markdown_path}")
    if arguments.render_pdf:
        render_pdf(docx_path, pdf_path, arguments.libreoffice)
    print(f"synthetic_environment_docx={docx_path}")
    print(f"artifact_fingerprint={sources.fingerprint}")
    if arguments.render_pdf:
        print(f"synthetic_environment_pdf={pdf_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
